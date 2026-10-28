
from pathlib import Path
import json

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports" / "phase14"
PRESENTATION = ROOT / "presentation"
NOTEBOOKS = ROOT / "notebooks"

for folder in (REPORTS, PRESENTATION, NOTEBOOKS):
    folder.mkdir(parents=True, exist_ok=True)

REPORT = r"""# FreshBasket Customer Churn Prediction
## Phase 14 — Final Project Report

**Program:** Tailwyndz Propel Internship  
**Assessment:** Customer Churn Prediction for a Retail Loyalty Program  
**Status:** Exploratory and provisional

## 1. Executive Summary

This project analyses FreshBasket customer profiles, monthly activity and churn labels. It builds customer-level features, compares classification models, evaluates held-out predictions, examines feature contributions and segments, and explores hypothetical retention scenarios.

The results are promising but are not ready for operational deployment. The Phase 11 audit identified unresolved inconsistencies between available activity records and churn labels. The strong recency signal also requires further investigation.

## 2. Business Objective

Identify customers who may stop purchasing so the loyalty team can investigate appropriate retention actions. The project covers data quality, feature engineering, model comparison, evaluation, explainability, segmentation and retention scenarios.

## 3. Dataset and Target

The workbook contains customer profiles, monthly activity, churn labels, a glossary and a data dictionary.

The supplied target definition describes churn as no transactions during April–June 2024 despite earlier purchase history. The intended feature cutoff is March 31, 2024; April–June 2024 activity is excluded from the predictor window.

The cleaned source tables contain 2,600 profiles, 28,015 monthly activity rows and 2,600 churn labels. The customer feature table contains 2,600 unique customer IDs. The eligible modelling cohort contains 2,227 customers; 373 were excluded under the insufficient-history eligibility rule.

## 4. Data Quality

- Duplicate rows and repeated customer-month records were audited.
- Two conflicting customer-month groups were investigated; conflicting values were not arbitrarily selected.
- Missing demographic and store-tier fields remain.
- Negative spend and spend/transaction inconsistencies remain documented concerns.
- Missing activity rows are not automatically interpreted as zero transactions.
- Original churn labels were preserved; no automatic relabelling was performed.

## 5. Feature Engineering and Validation

Features were aggregated from January 2023 through March 2024. They include transaction and spend summaries, engagement and support indicators, early/late-period summaries, trends, observed-month counts and recency-related measures.

The current experiment uses a stratified random customer-level split: 1,781 customers in the training pool and 446 held-out test customers. Preprocessing was fitted on training data.

This is not a true future-time holdout, so generalization to a later business period has not been established.

## 6. Held-Out Evaluation

The existing Phase 10 report records 446 test customers: 133 actual churners and 313 non-churners.

| Metric | Reported result |
|---|---:|
| Accuracy | 0.9260 |
| Precision | 0.8731 |
| Recall | 0.8797 |
| F1-score | 0.8764 |
| Specificity | 0.9457 |
| Balanced accuracy | 0.9127 |
| ROC-AUC | 0.9628 |
| PR-AUC / average precision | 0.9383 |
| Matthews correlation coefficient | 0.8236 |

Confusion matrix: TN=296, FP=17, FN=16, TP=117.

These metrics are transcribed from the existing report, not recalculated by this script. Strong scores do not rule out target-related leakage or label-quality problems.

## 7. Explainability and Segments

The Phase 11 SHAP analysis ranks MONTHS_SINCE_LAST_OBSERVED_ACTIVITY as the largest mean absolute SHAP contributor. Late-period transactions, observed feature months, email opens and support-related features also rank highly.

These are associations in the current experiment, not causal effects. City, membership tier, store price tier and signup cohort comparisons are descriptive.

## 8. Ablation Study

Phase 12 reported:
- Behaviour-only PR-AUC: 0.9787
- All features PR-AUC: 0.9773
- Without recency PR-AUC: 0.9071
- Without recency, late-period and trend features PR-AUC: 0.8390
- Profile-only PR-AUC: 0.3436

These results are provisional because the experiment uses a single internal validation split and the relationship between recency features and target construction has not been conclusively resolved.

## 9. Calibration and Retention Scenarios

The Phase 13 calibration audit reports a mean predicted probability of 0.3289 versus an observed churn rate of 0.2982, an absolute gap of about 3.07 percentage points. The Brier score is 0.0590 and log loss is 0.2194.

Calibration varies across risk ranges. Predicted probabilities should not yet be treated as reliable operational probabilities.

Retention scenarios assume prevention rates of 5%, 10% and 20%. These are assumptions, not measured intervention effects. A randomized retention test is needed to estimate incremental impact.

## 10. Critical Audit Findings and Limitations

- Five churn-positive labels have no earlier positive-transaction months in the available activity extract.
- A broader check found 125 non-churned customers with zero positive post-window months. This includes customers with no post-window activity rows and customers whose observed rows have no positive transactions.
- These are unresolved evidence gaps, not automatic proof that the labels are wrong. Original labels remain unchanged.
- MONTHS_SINCE_LAST_OBSERVED_ACTIVITY has a strong univariate AUC screen of approximately 0.9200. This does not by itself prove leakage, but the target definition and feature timing require verification.
- A date-boundary check alone does not prove that every predictor is independent of target construction.
- Random splitting does not establish performance on a genuinely later cohort.
- Segment differences are descriptive, not causal.

## 11. Recommendations

1. Clarify churn-label construction and missing-month semantics with the dataset owner.
2. Reconcile label/activity discrepancies against the source without silently changing labels.
3. Verify that all predictors use only permitted pre-label-window information.
4. Repeat evaluation using a defensible time-based holdout after resolving label and feature timing.
5. Repeat model comparison, calibration, SHAP and ablation on the revised setup.
6. Run a randomized retention pilot before estimating incremental benefit or deploying risk-based targeting.
7. Document package versions, preprocessing, random seeds and reproducibility instructions.

## 12. Conclusion

The project establishes an exploratory workflow and produces model comparisons, customer-level predictions, feature attribution, segment summaries, ablation results and illustrative retention scenarios.

Behavioural features are promising candidates for further validation, while profile-only features showed limited discrimination in this experiment. However, unresolved label/activity discrepancies and recency-related concerns mean that the current scores and feature rankings remain provisional. The next milestone is resolving data definitions and conducting leakage-aware future-period validation—not operational deployment.

## Appendix: Existing Evidence

- EDA: reports/phase3/PHASE3_EDA_FINDINGS.md
- Join validation: reports/phase4/phase4_validation_report.txt
- Feature validation: reports/phase5/phase5_feature_validation_report.txt
- Model comparison: reports/phase8/phase8_model_comparison_report.txt
- Held-out predictions: reports/phase9/phase9_prediction_report.txt
- Evaluation: reports/phase10/phase10_evaluation_report.txt
- Label/leakage audit: reports/phase11/phase11_final_audit_report.txt
- Ablation: reports/phase12/phase12_ablation_report.txt
- Retention/calibration: reports/phase13/phase13_retention_report.txt and phase13_calibration_report.txt
"""


def create_pdf(text, path):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ProjectTitle", parent=styles["Title"],
        fontSize=19, leading=23, spaceAfter=10
    ))
    styles.add(ParagraphStyle(
        name="Section", parent=styles["Heading1"],
        fontSize=14, leading=17, textColor=colors.HexColor("#174a6e"),
        spaceBefore=10
    ))
    styles.add(ParagraphStyle(
        name="SmallBody", parent=styles["BodyText"],
        fontSize=8.5, leading=11, spaceAfter=5
    ))

    story = []
    lines = text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for row in block:
                cells = [c.strip() for c in row.strip("|").split("|")]
                if all(set(c) <= set("-: ") for c in cells):
                    continue
                rows.append([
                    Paragraph(escape(c), styles["SmallBody"])
                    for c in cells
                ])
            if rows:
                table = Table(rows, repeatRows=1)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceaf2")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.extend([table, Spacer(1, 6)])
            continue

        clean = escape(line.replace("**", "").replace("`", ""))
        if line.startswith("# "):
            story.append(Paragraph(escape(line[2:]), styles["ProjectTitle"]))
        elif line.startswith("## "):
            story.append(Paragraph(escape(line[3:]), styles["Section"]))
        elif line.startswith("- "):
            story.append(Paragraph("&#8226; " + clean[2:], styles["SmallBody"]))
        else:
            story.append(Paragraph(clean, styles["SmallBody"]))
        i += 1

    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        rightMargin=17*mm, leftMargin=17*mm,
        topMargin=16*mm, bottomMargin=16*mm
    )
    doc.build(story)


def create_presentation(path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = [
        ("Customer Churn Prediction", [
            "FreshBasket Retail Loyalty Program",
            "Tailwyndz Propel Internship — Assessment 4",
            "Final presentation: exploratory and provisional"
        ]),
        ("Business Objective", [
            "Identify customers who may stop purchasing.",
            "Use pre-cutoff behaviour to rank risk.",
            "Inform retention investigation—not automate decisions yet."
        ]),
        ("Dataset and Target", [
            "2,600 profiles; 28,015 cleaned activity rows; 2,600 churn labels.",
            "Target window: April–June 2024; feature cutoff: March 31, 2024.",
            "Eligible cohort: 2,227 customers; 373 excluded by history rule."
        ]),
        ("Data Quality", [
            "Duplicates and repeated customer-month records were audited.",
            "Conflicting customer-month records were not arbitrarily resolved.",
            "Missing activity rows are not automatically zero transactions.",
            "Label/activity evidence gaps remain unresolved; original labels preserved."
        ]),
        ("Feature Engineering", [
            "Pre-cutoff transaction, spend, engagement and support summaries.",
            "Early/late activity, trends, observed-month and recency measures.",
            "Recency and missingness semantics need further review."
        ]),
        ("Validation Design", [
            "Stratified random split: 1,781 training-pool / 446 held-out test customers.",
            "Preprocessing fitted on training data.",
            "Not a future-time holdout; future generalization is unproven."
        ]),
        ("Held-Out Performance", [
            "Accuracy 0.9260 | Precision 0.8731 | Recall 0.8797 | F1 0.8764",
            "ROC-AUC 0.9628 | PR-AUC 0.9383 | Balanced accuracy 0.9127",
            "Confusion matrix: TN=296, FP=17, FN=16, TP=117.",
            "Values transcribed from Phase 10 report, not recomputed here."
        ]),
        ("Explainability", [
            "Largest mean absolute SHAP contributor: MONTHS_SINCE_LAST_OBSERVED_ACTIVITY.",
            "Late-period transactions, observed months, email opens and support also rank highly.",
            "Feature attribution is association, not causation or proof of leakage-free validity."
        ]),
        ("Ablation Study", [
            "Behaviour-only PR-AUC 0.9787; all features 0.9773.",
            "Without recency PR-AUC 0.9071.",
            "Without recency plus late/trend features PR-AUC 0.8390.",
            "Profile-only PR-AUC 0.3436; rankings remain provisional."
        ]),
        ("Calibration and Scenarios", [
            "Mean predicted probability 0.3289 vs observed rate 0.2982.",
            "Brier score 0.0590; log loss 0.2194; calibration varies by risk range.",
            "5%, 10%, 20% prevention rates are assumptions, not measured effects."
        ]),
        ("Critical Limitations", [
            "Five churn-positive labels lack earlier positive-transaction months in the extract.",
            "125 non-churned customers have zero positive post-window months.",
            "Recency feature univariate AUC screen is about 0.9200.",
            "Random split does not prove later-period generalization."
        ]),
        ("Recommendations", [
            "Clarify label construction and missing-month semantics with the data owner.",
            "Reconcile evidence gaps without silently changing labels.",
            "Verify feature timing and rerun evaluation with a time-based holdout.",
            "Repeat calibration and interpretation; run a randomized retention pilot."
        ]),
        ("Conclusion and Q&A", [
            "Behavioural features are promising candidates for further validation.",
            "Profile-only features showed limited discrimination in this experiment.",
            "Scores, probabilities and retention scenarios are provisional.",
            "Next: resolve data definitions and validate on a future period."
        ]),
    ]

    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(247, 250, 252)

        box = slide.shapes.add_textbox(
            Inches(.65), Inches(.45), Inches(12), Inches(.8)
        )
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(23, 74, 110)

        box = slide.shapes.add_textbox(
            Inches(.85), Inches(1.55), Inches(11.7), Inches(5.2)
        )
        tf = box.text_frame
        tf.word_wrap = True
        for index, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(20 if len(bullet) < 105 else 17)
            p.font.color.rgb = RGBColor(38, 50, 56)
            p.space_after = Pt(14)

        footer = slide.shapes.add_textbox(
            Inches(.7), Inches(7.12), Inches(12), Inches(.2)
        )
        fp = footer.text_frame.paragraphs[0]
        fp.text = "TAILWYNDZ PROPEL | Exploratory findings — not deployment-ready"
        fp.font.size = Pt(8)
        fp.font.color.rgb = RGBColor(100, 110, 118)

    prs.save(path)


def create_eda_notebook():
    path = NOTEBOOKS / "02_FreshBasket_EDA.ipynb"
    if path.exists():
        print("Preserving existing notebook:", path.relative_to(ROOT))
        return

    cells = [
        {
            "cell_type": "markdown", "metadata": {},
            "source": [
                "# FreshBasket Loyalty Churn — EDA\n",
                "Reproducible workbook inspection, missingness and duplicate checks. ",
                "This notebook does not modify source data or churn labels."
            ]
        },
        {
            "cell_type": "code", "execution_count": None, "metadata": {},
            "outputs": [],
            "source": [
                "from pathlib import Path\n",
                "import pandas as pd\n",
                "import matplotlib.pyplot as plt\n",
                "ROOT = Path.cwd()\n",
                "if ROOT.name.lower() != 'tailwyndz_churn_prediction' and ROOT.parent.name.lower() == 'tailwyndz_churn_prediction': ROOT = ROOT.parent\n",
                "WORKBOOK = ROOT / 'data' / 'raw' / 'FreshBasket_Loyalty_Churn_Dataset.xlsx'\n",
                "assert WORKBOOK.exists(), f'Workbook not found: {WORKBOOK}'\n",
                "xls = pd.ExcelFile(WORKBOOK)\n",
                "print('Sheets:', xls.sheet_names)\n",
                "tables = {s: pd.read_excel(WORKBOOK, sheet_name=s) for s in xls.sheet_names}\n"
            ]
        },
        {
            "cell_type": "code", "execution_count": None, "metadata": {},
            "outputs": [],
            "source": [
                "for name, df in tables.items():\n",
                "    print('\\\\n', name, 'shape=', df.shape, 'exact duplicates=', int(df.duplicated().sum()))\n",
                "    display(df.head())\n",
                "    display(df.isna().sum().sort_values(ascending=False).head(20).to_frame('missing_count'))\n"
            ]
        },
        {
            "cell_type": "code", "execution_count": None, "metadata": {},
            "outputs": [],
            "source": [
                "for name, df in tables.items():\n",
                "    if 'CHURNED' in df.columns:\n",
                "        print('Label table:', name)\n",
                "        display(df['CHURNED'].value_counts(dropna=False).rename_axis('CHURNED').to_frame('rows'))\n",
                "        ax = df['CHURNED'].value_counts().sort_index().plot(kind='bar', rot=0)\n",
                "        ax.set_title('Churn label distribution'); ax.set_xlabel('CHURNED'); ax.set_ylabel('Rows')\n",
                "        plt.tight_layout(); plt.show()\n"
            ]
        },
        {
            "cell_type": "markdown", "metadata": {},
            "source": [
                "## Interpretation guardrails\n",
                "- Confirm the target definition before modelling.\n",
                "- Keep April–June 2024 activity out of predictors if that is the target window.\n",
                "- Missing activity months do not automatically mean zero transactions.\n",
                "- Preserve conflicting customer-month records pending clarification.\n",
                "- Use future-time validation after label and feature timing are verified."
            ]
        }
    ]
    notebook = {
        "cells": cells,
        "metadata": {
            "kernelspec": {
                "display_name": "Python 3",
                "language": "python",
                "name": "python3"
            },
            "language_info": {"name": "python"}
        },
        "nbformat": 4,
        "nbformat_minor": 5
    }
    path.write_text(json.dumps(notebook, indent=2), encoding="utf-8")
    print("Created:", path.relative_to(ROOT))


def main():
    md_path = REPORTS / "final_report.md"
    md_path.write_text(REPORT, encoding="utf-8")
    print("Created:", md_path.relative_to(ROOT))

    create_eda_notebook()

    try:
        create_pdf(REPORT, REPORTS / "final_report.pdf")
        print("Created: reports/phase14/final_report.pdf")
    except ImportError:
        print("PDF dependency missing. Run: python -m pip install reportlab")

    try:
        create_presentation(PRESENTATION / "final_presentation.pptx")
        print("Created: presentation/final_presentation.pptx")
    except ImportError:
        print("PPTX dependency missing. Run: python -m pip install python-pptx")

    print("\nPhase 14 generation complete. Review every output before submission.")


if __name__ == "__main__":
    main()
